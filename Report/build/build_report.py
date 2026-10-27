# -*- coding: utf-8 -*-
"""Build Report/Report.docx from the Report folder (diagrams in Images/, screenshots in In-App images/).
Run:  python build/build_report.py   then   powershell -File build/update_fields.ps1   (TOC / list of figures / tables + PDF)."""
import csv, json, os, re, sys
from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image

HERE = os.path.dirname(os.path.abspath(__file__))
REPORT = os.path.dirname(HERE)
sys.path.insert(0, HERE)
import report_text as T
import report_data as D
from narrative import read as read_narrative
from sqlparse import queries, schema, TABLE_ORDER

IMG = os.path.join(REPORT, "Images")
SHOT = os.path.join(REPORT, "In-App images")
DATA = os.path.join(HERE, "data")
OUT = os.path.join(REPORT, "Report.docx")
FONT, MONO = "TH Sarabun New", "Consolas"
UC_NAMES = {"UC01": "Create Project Profile", "UC02": "Store Project Solution Details", "UC03": "Check Project Details",
            "UC04": "Assign Engineer to Project", "UC05": "Set Project Test Calendar Notification",
            "UC06": "Add Customer New Requirement", "UC07": "Approve New Project Requirement", "UC08": "Close Project"}


# ------------------------------------------------------------------ low-level helpers
def set_run_font(run, name=FONT, size=None, bold=None, cs_name=None):
    rpr = run._element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts"); rpr.insert(0, rf)
    for a in ("w:ascii", "w:hAnsi", "w:eastAsia"):
        rf.set(qn(a), name)
    rf.set(qn("w:cs"), cs_name or FONT)
    if size:
        run.font.size = Pt(size)
        szcs = rpr.find(qn("w:szCs"))
        if szcs is None:
            szcs = OxmlElement("w:szCs"); rpr.append(szcs)
        szcs.set(qn("w:val"), str(int(size * 2 if name == FONT else max(size * 2, 26))))
    if bold is not None:
        run.bold = bold
        b = rpr.find(qn("w:bCs"))
        if bold and b is None:
            rpr.append(OxmlElement("w:bCs"))


def style_font(style, size, bold=False, color=None):
    style.font.name = FONT
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.italic = False
    if color is not None:
        style.font.color.rgb = color
    rpr = style.element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts"); rpr.insert(0, rf)
    for a in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rf.set(qn(a), FONT)
    for a in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        if rf.get(qn(a)) is not None:
            del rf.attrib[qn(a)]
    szcs = rpr.find(qn("w:szCs"))
    if szcs is None:
        szcs = OxmlElement("w:szCs"); rpr.append(szcs)
    szcs.set(qn("w:val"), str(size * 2))
    if bold and rpr.find(qn("w:bCs")) is None:
        rpr.append(OxmlElement("w:bCs"))


def add_field(par, instr, placeholder="", bold=False, size=None):
    def r(el):
        run = par.add_run(); set_run_font(run, size=size, bold=bold); run._r.append(el); return run
    b = OxmlElement("w:fldChar"); b.set(qn("w:fldCharType"), "begin"); r(b)
    it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = instr; r(it)
    s = OxmlElement("w:fldChar"); s.set(qn("w:fldCharType"), "separate"); r(s)
    ph = par.add_run(placeholder); set_run_font(ph, size=size, bold=bold)
    e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), "end"); r(e)


def shade(cell, fill):
    tcpr = cell._tc.get_or_add_tcPr()
    sh = OxmlElement("w:shd"); sh.set(qn("w:val"), "clear"); sh.set(qn("w:color"), "auto"); sh.set(qn("w:fill"), fill)
    tcpr.append(sh)


def rich(par, text, size=None, bold=False, mono=False):
    """[[...]] = yellow highlight (placeholder / assumption), **...** = bold."""
    for part in re.split(r"(\[\[.*?\]\]|\*\*.*?\*\*)", text):
        if not part:
            continue
        hl = part.startswith("[[")
        bb = part.startswith("**")
        txt = part[2:-2] if (hl or bb) else part
        run = par.add_run(txt)
        set_run_font(run, name=MONO if mono else FONT, size=size, bold=bold or bb)
        if hl:
            run.font.highlight_color = WD_COLOR_INDEX.YELLOW
    return par


# ------------------------------------------------------------------ document wrapper
class Rep:
    def __init__(self):
        self.doc = Document()
        self.landscape = False
        self.fresh_page = True
        self.pending_break = False
        d = self.doc
        st = d.styles
        style_font(st["Normal"], 16)
        st["Normal"].paragraph_format.space_after = Pt(4)
        st["Normal"].paragraph_format.line_spacing = 1.0
        for lvl, size in ((1, 22), (2, 18), (3, 16)):
            h = st[f"Heading {lvl}"]
            style_font(h, size, True, RGBColor(0, 0, 0))
            h.paragraph_format.space_before = Pt(12 if lvl > 1 else 0)
            h.paragraph_format.space_after = Pt(6)
            h.paragraph_format.keep_with_next = True
        st["Heading 1"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        style_font(st["Caption"], 16, False, RGBColor(0, 0, 0))
        st["Caption"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        st["Caption"].paragraph_format.space_after = Pt(10)
        style_font(st["List Bullet"], 16)
        for name in ("TOC 1", "TOC 2", "TOC 3", "Table of Figures"):
            try:
                s = st[name]
            except KeyError:
                s = st.add_style(name, 1)
            style_font(s, 16, name == "TOC 1")
        sec = d.sections[0]
        self._page(sec, False)
        sec.different_first_page_header_footer = True
        fp = sec.footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_field(fp, " PAGE ", "1", size=14)

    @staticmethod
    def _page(sec, landscape):
        if landscape:
            sec.orientation = WD_ORIENT.LANDSCAPE
            sec.page_width, sec.page_height = Cm(29.7), Cm(21.0)
        else:
            sec.orientation = WD_ORIENT.PORTRAIT
            sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
        sec.left_margin, sec.right_margin = Cm(3.0), Cm(2.0)
        sec.top_margin, sec.bottom_margin = Cm(2.5), Cm(2.0)

    @property
    def width(self):
        return 24.7 if self.landscape else 16.0

    @property
    def height(self):
        return 15.0 if self.landscape else 23.5

    def orient(self, landscape):
        if landscape != self.landscape:
            self.pending_break = False
            sec = self.doc.add_section(WD_SECTION.NEW_PAGE)
            self._page(sec, landscape)
            self.landscape = landscape
            self.fresh_page = True

    def page_break(self):
        """the next paragraph starts a new page (page-break-before, so no empty spill-over page)"""
        if not self.fresh_page:
            self.pending_break = True
            self.fresh_page = True

    def para(self, style=None):
        par = self.doc.add_paragraph(style=style) if style else self.doc.add_paragraph()
        self._apply_break(par)
        return par

    def _apply_break(self, par):
        if self.pending_break:
            par.paragraph_format.page_break_before = True
            self.pending_break = False

    def h(self, lvl, text, new_page=False):
        if new_page:
            self.page_break()
        p = self.doc.add_heading(level=lvl)
        self._apply_break(p)
        rich(p, text)
        self.fresh_page = False
        return p

    def p(self, text, indent=True, size=None, align=None, bold=False):
        par = self.para()
        if indent:
            par.paragraph_format.first_line_indent = Cm(1.25)
        if align == "center":
            par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif align == "justify" or align is None:
            par.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY if indent else None
        rich(par, text, size=size, bold=bold)
        self.fresh_page = False
        return par

    def bullet(self, text, size=None):
        par = self.para("List Bullet")
        rich(par, text, size=size)
        self.fresh_page = False
        return par

    def caption(self, kind, text):
        par = self.para("Caption")
        label = "รูปที่ " if kind == "Figure" else "ตารางที่ "
        r = par.add_run(label); set_run_font(r, bold=True)
        add_field(par, f" SEQ {kind} \\* ARABIC ", "1", bold=True)
        r = par.add_run(" " + text); set_run_font(r)
        if kind == "Table":
            par.paragraph_format.keep_with_next = True
            par.paragraph_format.space_after = Pt(2)
        self.fresh_page = False
        return par

    def figure(self, path, caption, max_w=None, max_h=None):
        w, h = Image.open(path).size
        mw, mh = max_w or self.width, max_h or self.height
        scale = min(mw / w, mh / h)
        par = self.para()
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        par.paragraph_format.keep_with_next = True
        par.add_run().add_picture(path, width=Cm(w * scale))
        self.caption("Figure", caption)

    def table(self, headers, rows, widths, size=14, mono_cols=(), center_cols=(), header_fill="D9E2F3",
              first_col_fill=None, mono_size=8.5):
        t = self.doc.add_table(rows=0, cols=len(widths))
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        all_rows = ([headers] if headers else []) + list(rows)
        for i, row in enumerate(all_rows):
            cells = t.add_row().cells
            for j, val in enumerate(row):
                c = cells[j]
                c.width = Cm(widths[j])
                c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
                lines = str(val if val is not None else "").split("\n")
                is_head = headers and i == 0
                mono_cell = (j in mono_cols) and not is_head
                for k, line in enumerate(lines):
                    mono = mono_cell and not line.startswith("ผ่าน")
                    par = c.paragraphs[0] if k == 0 else c.add_paragraph()
                    par.paragraph_format.space_after = Pt(0)
                    if is_head or j in center_cols:
                        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    rich(par, line, size=mono_size if mono else size, bold=is_head, mono=mono)
                if is_head:
                    shade(c, header_fill)
                elif first_col_fill and j == 0:
                    shade(c, first_col_fill)
            if headers and i == 0:
                trpr = t.rows[0]._tr.get_or_add_trPr()
                th = OxmlElement("w:tblHeader"); th.set(qn("w:val"), "true"); trpr.append(th)
        # fixed layout so Word keeps the widths
        tblpr = t._tbl.tblPr
        lay = OxmlElement("w:tblLayout"); lay.set(qn("w:type"), "fixed"); tblpr.append(lay)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(2)
        self.fresh_page = False
        return t

    def save(self):
        self.doc.save(OUT)


def img(name):
    return os.path.join(IMG, name)


def shot(name):
    return os.path.join(SHOT, name)


# ------------------------------------------------------------------ sections
# Cover = Report/Cover template.docx, paragraph for paragraph: (text, Angsana New size in pt, centered).
# Blank lines keep the template's spacing; [[...]] placeholders sit on blank lines of the template.
COVER_FONT = "Angsana New"
COVER_LINES = [
    ("", 11, True),     # template default (Arial 11)
    ("Project Management Program", 15, True),
    ("", 14, True), ("", 14, True), ("", 14, True), ("", 14, True), ("", 15, True),
    ("จัดทำโดย", 15, True),
    ("", 15, True),
    ("นายณัฐพล ทวีสินธนานนท์ 6710450830", 15, True),
    ("[[ชื่อ-นามสกุล รหัสนักศึกษา (สมาชิกคนอื่น)]]", 15, True),
    ("", 15, True), ("", 15, False), ("", 15, True),
    ("เสนอ", 15, True),
    ("[[ชื่ออาจารย์ผู้สอน]]", 15, True), ("", 15, True), ("", 15, True), ("", 15, True), ("", 15, False), ("", 15, False),
    ("รายงานนี้เป็นส่วนหนึ่งของวิชา การวิเคราะห์และออกแบบระบบ", 15, True),
    ("ภาคเรียนที่ 1 ปีการศึกษา 2569", 15, True),
    ("มหาวิทยาลัยเกษตรศาสตร์ บางเขน", 15, True),
]


def cover(r):
    def para(center):
        p = r.para()
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.15
        if center:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        return p
    logo = para(True)
    logo.add_run().add_picture(img("KU_logo.png"), width=Cm(3.75))
    def lang_th(rpr):   # the template's default language is Thai; Word's line height for Angsana depends on it
        lg = OxmlElement("w:lang"); lg.set(qn("w:val"), "th"); rpr.append(lg)
    for text, size, center in COVER_LINES:
        p = para(center)
        if text:
            rich(p, text, size=size)
            for run in p.runs:
                set_run_font(run, name=COVER_FONT, size=size, cs_name=COVER_FONT)
                lang_th(run._element.get_or_add_rPr())
        else:   # empty line: the paragraph mark carries the template's font / size, so the line height matches
            rpr = OxmlElement("w:rPr")
            rf = OxmlElement("w:rFonts")
            for k in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
                rf.set(qn(k), "Arial" if size == 11 else COVER_FONT)
            rpr.append(rf)
            for k in ("w:sz", "w:szCs"):
                e = OxmlElement(k); e.set(qn("w:val"), str(size * 2)); rpr.append(e)
            lang_th(rpr)
            p._p.get_or_add_pPr().append(rpr)
    r.fresh_page = False


def front(r):
    for title, instr in (("สารบัญ", ' TOC \\o "1-2" \\h \\z \\u '), ("สารบัญรูป", ' TOC \\h \\z \\c "Figure" '),
                         ("สารบัญตาราง", ' TOC \\h \\z \\c "Table" ')):
        r.page_break()
        p = r.para(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        rich(p, title, size=22, bold=True)
        par = r.para()
        add_field(par, instr, "(คลิกขวา → Update Field เพื่อแสดงเลขหน้า)")
        r.fresh_page = False
    r.page_break()
    p = r.para(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rich(p, "บทคัดย่อ", size=22, bold=True)
    for t in T.ABSTRACT_TH:
        r.p(t, size=14)
    r.p(T.KEYWORDS_TH, indent=False, size=14, bold=True)
    p = r.para(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rich(p, "Abstract", size=22, bold=True)
    for t in T.ABSTRACT_EN:
        r.p(t, size=13)
    r.p(T.KEYWORDS_EN, indent=False, size=13, bold=True)


def ch1(r):
    r.h(1, "บทที่ 1 บทนำ", new_page=True)
    r.h(2, "1.1 ความเป็นมาและหน่วยงาน")
    for t in T.INTRO:
        r.p(t)
    r.h(2, "1.2 ปัญหาที่ Application จะเข้าไปแก้ไข")
    r.p(T.PROBLEMS_INTRO)
    r.caption("Table", "ปัญหาของหน่วยงานพร้อมตัวเลขสนับสนุน")
    r.table(("ที่", "ปัญหา", "ตัวเลขสนับสนุน"), [(str(i + 1), a, b) for i, (a, b) in enumerate(T.PROBLEMS)], (1.0, 5.0, 10.0),
            center_cols=(0,))
    r.h(2, "1.3 ขอบเขตของระบบ")
    r.p("ระบบรองรับผู้ใช้ 4 กลุ่ม คือ Sale, Presale Engineer, PM และ Technical Engineer มี 8 Use Case (UC01–UC08) และฟังก์ชันพื้นฐาน "
        "สมัครสมาชิก (F1) เข้าสู่ระบบ (F2) โปรไฟล์ (F3) และกระดิ่งแจ้งเตือน (F4) งานที่ยังเป็น manual ได้แก่ การจัดซื้อ / ตรวจ Stock อุปกรณ์ "
        "การ Outsource การติดตั้งที่ Site และการจัดทำเอกสาร UAT / Handover ระบบไม่ส่งอีเมลจริง (บันทึก log แทน) แต่แจ้งเตือนในระบบทุกขั้นตอน")


def ch2(r):
    r.h(1, "บทที่ 2 Business Process", new_page=True)
    r.h(2, "2.1 Business Process เดิม (As-is)")
    r.p("Business Process เดิมแบ่งเป็น 5 Swim lane ตามผู้ปฏิบัติงาน ทุก Process มีเลขกำกับตามลำดับการทำงานจากบนลงล่าง ซ้ายไปขวา "
        "กล่องสีแดงคือ Process ที่เกิด Pain Point ระบบเดิมทำงานโดยใช้อีเมลและเอกสาร (physical) เป็นหลัก Pain Point ทั้ง 5 ข้อ ได้แก่")
    for t in ("Pain Point 1 – การนัด Meeting หรือ Test กับลูกค้าคลาดเคลื่อน (Process 17)",
              "Pain Point 2 – สถานที่ทำงานไม่ตรงตามนัด (Process 14)",
              "Pain Point 3 – Project delay หากมีการเพิ่ม Requirement ระหว่าง Project (Process 18)",
              "Pain Point 4 – PM ไม่รู้ความสามารถที่แน่นอนของ Engineer เพื่อจัดวางให้เหมาะกับแต่ละ Project (Process 6)",
              "Pain Point 5 – มีปัญหาเรื่องรายละเอียด Project ในกรณีส่งงานต่อกัน (Process 4)"):
        r.bullet(t)
    r.figure(img("As-If.png"), "Business Process เดิม (As-is)")
    r.h(2, "2.2 คำอธิบาย Business Process เดิม")
    r.p("ระบบเดิมทำงานด้วยอีเมลและเอกสารกระดาษเป็นหลัก แต่ละ Swim lane ทำงานดังนี้ (ตัวเลขในวงเล็บคือเลข Process)")
    for lane, txt in T.ASIS_DESC:
        r.p(f"**{lane}** – {txt}")
    r.h(2, "2.3 ปัญหาของระบบเดิม")
    r.p("ปัญหาหลักของระบบเดิมคำนวณขนาดความเสียหายได้ดังตารางต่อไปนี้ (ค่าที่ไฮไลต์เป็นสมมติฐานจากการสัมภาษณ์ ให้แทนด้วยข้อมูลจริง) "
        "ทุกปัญหาเป็นปัญหาที่ระบบใหม่แก้ได้ และวัดผลได้ว่าลดลง")
    r.caption("Table", "ปัญหาของระบบเดิม ขนาดความเสียหาย และตัวชี้วัดหลังใช้ระบบใหม่")
    r.table(("Pain Point (Process)", "การคำนวณความเสียหาย / เวลาที่เสียไป", "ระบบใหม่แก้อย่างไร / วัดผลอย่างไร"), T.ASIS_PROBLEMS,
            (4.2, 5.6, 6.2))
    r.p("รวมแล้วระบบเดิมเสียเวลาทำงานประมาณ [[288 + 108 + 240 = 636 ชั่วโมง-คน/ปี]] ค่าเดินทางซ้ำ [[18,000 บาท/ปี]] และโครงการล่าช้าเฉลี่ย "
        "[[10 วันทำการต่อโครงการ]] จากการรออนุมัติ Requirement")
    r.h(2, "2.4 Business Process ใหม่ (To-be)")
    r.p("Business Process ใหม่ใช้ Swim lane ชุดเดิม (ไม่มี Swim lane ของระบบ) กล่องสีม่วงคือ Process ที่ปรับปรุงโดยใช้ Project Management Program "
        "พร้อมเลข Use Case กล่องสีขาวคือ Process ที่ยังเป็น manual จัดวางให้เส้นตัดกันน้อยที่สุด และใช้จุดเชื่อม A แทนเส้นที่ย้อนกลับขึ้นด้านบน")
    r.figure(img("To-Be.png"), "Business Process ใหม่ (To-be)")
    r.h(2, "2.5 คำอธิบาย Business Process ใหม่")
    for lane, txt in T.TOBE_DESC:
        r.p(f"**{lane}** – {txt}")
    r.p(T.TOBE_MANUAL)
    r.caption("Table", "Process ใน To-be ที่แก้ปัญหาของระบบเดิม")
    r.table(("ปัญหาของระบบเดิม", "Process ที่แก้ปัญหา"), T.TOBE_FIX, (5.5, 10.5))
    r.h(2, "2.6 Business Process ที่สัมพันธ์กับ Use Case")
    r.p("รูปต่อไปนี้วง Use Case ลงบน To-be BP วงเส้นทึบสีม่วงคือ Use Case ที่ใช้คอมพิวเตอร์ (UC01–UC08) วงเส้นประสีส้มคือ Use Case ที่ยังเป็น manual (M1–M7) ดังนี้")
    for t in ("M1 จัดหา Project (Sale) – Process 1", "M2 ออกแบบ Solution และวางแผน Project (Presale Engineer) – Process 4",
              "M3 เริ่ม Project และแจ้งอุปกรณ์ที่ต้องใช้ (Presale Engineer) – Process 5, 11",
              "M4 จัดซื้อ / ตรวจ Stock อุปกรณ์ (Sale, จัดซื้อ / Stock) – Process 12, 13, 14",
              "M5 วางแผนงาน Project (PM) – Process 8, 10", "M6 ติดตั้งและทดสอบระบบภายใน (Technical Engineer) – Process 16, 17, 18",
              "M7 จัดทำเอกสาร UAT และ Handover (Technical Engineer) – Process 24"):
        r.bullet(t)
    r.figure(img("BP-UseCase.png"), "Business Process ที่สัมพันธ์กับ Use Case")
    r.h(2, "2.7 ตารางจับคู่ Business Process กับ Use Case")
    r.caption("Table", "Use Case ที่ได้จาก Business Process (ทุก UC มาจาก BP)")
    r.table(("UC", "ชื่อ Use Case", "Actor (Swim lane)", "Process ใน To-be", "Pain Point"), T.UC_FROM_BP,
            (1.4, 4.6, 4.6, 2.6, 2.8), center_cols=(0, 3, 4))
    r.caption("Table", "Process ทั้งหมดใน To-be BP กับ Use Case (คอมพิวเตอร์ / manual)")
    r.table(("Process", "ชื่อ Process", "Swim lane", "Use Case"), T.BP_UC_ROWS, (1.6, 6.0, 3.6, 4.8), center_cols=(0,), size=13)


def ch3(r):
    r.h(1, "บทที่ 3 Use Case Diagram และ CRUD Table", new_page=True)
    r.h(2, "3.1 Use Case Diagram (ฉบับแรก)")
    r.p(T.UC_FIRST_DESC)
    r.figure(img("Use-Case_Diagram_First_.png"), "Use Case Diagram ฉบับแรก (7 Use Case)", max_h=15)
    r.h(2, "3.2 CRUD Table (ฉบับแรก)")
    r.p(T.NOUNS)
    r.figure(img("CRUD_Table_First_.png"), "CRUD Table ฉบับแรก (UC01–UC07 × Entity)")
    r.h(2, "3.3 คำอธิบาย CRUD Table และ Use Case ที่ยังขาด")
    r.p(T.CRUD_FIRST_ANALYSIS[0])
    for t in T.CRUD_FIRST_ANALYSIS[1:]:
        r.bullet(t)
    r.h(2, "3.4 Use Case Diagram ที่ปรับปรุงแล้ว")
    r.p(T.UC_REVISED_DESC)
    r.figure(img("Use-Case_Diagram.png"), "Use Case Diagram ที่ปรับปรุงแล้ว (8 Use Case)", max_h=15)
    r.h(2, "3.5 CRUD Table ฉบับทบทวน")
    r.figure(img("CRUD_Table.png"), "CRUD Table ฉบับทบทวน (UC01–UC08 × Entity)")
    for t in T.CRUD_REVISED_DESC:
        r.p(t)
    r.caption("Table", "CRUD ของฟังก์ชันพื้นฐาน F1–F4 (ไม่ได้มาจาก Process ใน BP)")
    heads = ("ฟังก์ชัน",) + tuple(t.replace("_", "_ ") for t in TABLE_ORDER)
    r.table(heads, T.CRUD_F_ROWS, (2.6,) + (1.2,) * 11, size=9, center_cols=tuple(range(1, 12)))


def steps_of(cell):
    """split a narrative cell into {step no: text}"""
    out, cur = {}, None
    for line in cell.split("\n"):
        m = re.match(r"\s*Step (\d+):\s*(.*)", line)
        if m:
            cur = int(m.group(1)); out[cur] = [m.group(2)]
        elif cur is not None:
            out[cur].append(line)
    return {k: "\n".join(v).strip("\n") for k, v in out.items()}


def mark_sql(text):
    """lines after 'Query N.M:' up to a blank line are SQL -> wrap in a mono marker"""
    res, in_sql = [], False
    for line in text.split("\n"):
        if re.match(r"\s*Query \d+\.\d+", line):
            in_sql = True; res.append(("**" + line.strip() + "**", False)); continue
        if not line.strip():
            in_sql = False; continue
        res.append((line, in_sql))
    return res


def uc_description(r, uc, nar):
    rows = nar["rows"]
    head = {row[0]: row[1] for row in rows if len(row) >= 2 and row[0] in
            ("Main Actor", "Describe", "Pre-Condition", "Post-Condition", "Alternative / Exception Flow")}
    actor_cell = system_cell = ""
    for row in rows:
        if len(row) >= 2 and "Step" in row[0] + row[1] and row[0] not in head:
            actor_cell, system_cell = row[0], row[1]
    actor = head.get("Main Actor", "").replace("Tech Engineer", "Technical Engineer")
    r.caption("Table", f"Use Case Description {uc} {UC_NAMES[uc]}")
    t = r.doc.add_table(rows=0, cols=2)
    t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    W = (6.0, 10.0)

    def row2(a, b, fill=None, span=False, a_bold=True):
        cells = t.add_row().cells
        if span:
            m = cells[0].merge(cells[1])
            rich(m.paragraphs[0], a, size=14, bold=True)
            if fill:
                shade(m, fill)
            return
        for c, w in zip(cells, W):
            c.width = Cm(w)
        rich(cells[0].paragraphs[0], a, size=14, bold=a_bold)
        lines = b.split("\n")
        for k, line in enumerate(lines):
            par = cells[1].paragraphs[0] if k == 0 else cells[1].add_paragraph()
            par.paragraph_format.space_after = Pt(0)
            rich(par, line, size=14)
        if fill:
            shade(cells[0], fill)
    row2(f"Use Case: {uc} {UC_NAMES[uc]}", "", "D9E2F3", span=True)
    for k, label in (("Main Actor", "Actor"), ("Describe", "Description"), ("Pre-Condition", "Pre-Condition"),
                     ("Post-Condition", "Post-Condition")):
        row2(label, actor if k == "Main Actor" else head.get(k, ""), "F2F2F2")
    hdr = t.add_row().cells
    for c, txt in zip(hdr, ("Actor Action", "System Response (Query)")):
        rich(c.paragraphs[0], txt, size=14, bold=True); shade(c, "D9E2F3")
        c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    a, s = steps_of(actor_cell), steps_of(system_cell)
    for n in sorted(set(a) | set(s)):
        cells = t.add_row().cells
        for c, w in zip(cells, W):
            c.width = Cm(w)
        target = cells[0] if n in a else cells[1]
        text = a.get(n) or s.get(n)
        first = True
        for line, is_sql in mark_sql(f"Step {n}: " + text):
            par = target.paragraphs[0] if first else target.add_paragraph()
            first = False
            par.paragraph_format.space_after = Pt(0)
            rich(par, line, size=8.5 if is_sql else 14, mono=is_sql)
    row2("Alternative / Exception Flow", head.get("Alternative / Exception Flow", ""), "F2F2F2")
    tblpr = t._tbl.tblPr
    lay = OxmlElement("w:tblLayout"); lay.set(qn("w:type"), "fixed"); tblpr.append(lay)
    r.para()
    r.fresh_page = False


def ch4(r):
    r.h(1, "บทที่ 4 Use Case Description, Sequence และ Collaboration Diagram", new_page=True)
    r.p("แต่ละ Use Case นำเสนอเป็นชุด คือ Use Case Description → Sequence Diagram → Collaboration Diagram เลข Step ใน Description, "
        "เลข Query (N.M) และเลข Message ใน Diagram ใช้ระบบเดียวกัน ได้แก่")
    r.bullet("**N.** = Step N ของ Use Case Description")
    r.bullet("**N.M** = Query N.M ที่ระบบรันใน Step N (คำสั่ง SQL เดียวกับใน Description และตาราง Query)")
    r.bullet("**Na, Nb …** = Message ภายในระหว่าง Boundary → Controller → API ที่นำไปสู่ Step N (เช่น 2a onClick(), 2b GET /api/…)")
    r.p("Object ที่ติดกับ Actor เป็น <<Boundary>> (Button / Form / Page ซึ่งเป็นหน้าจอจริงในโปรแกรม) ถัดไปเป็น <<Controller>> (ฟังก์ชัน JavaScript "
        "และ DetectBlank / File Check ที่ตรวจข้อมูล) <<API>> (REST endpoint ใน Spring Boot) และ <<Entity>> (ตารางในฐานข้อมูล) "
        "<<API>> เชื่อมต่อฐานข้อมูลผ่าน NamedParameterJdbcTemplate (database connect ของโปรแกรม) Sequence Diagram แสดงเฉพาะทางหลัก (main path) "
        "ส่วนทางเลือกอยู่ใน Alternative / Exception Flow ของ Description Collaboration Diagram แปลงมาจาก Sequence Diagram โดยใช้ Object ชุดเดียวกัน "
        "และ Message มีเลขกำกับตรงกันทุกตัว (ตรวจสอบด้วยสคริปต์แล้ว)")
    nar = read_narrative("with")
    for i, uc in enumerate(UC_NAMES, 1):
        r.h(2, f"4.{i} {uc} {UC_NAMES[uc]}", new_page=i > 1)
        r.h(3, f"4.{i}.1 Use Case Description")
        uc_description(r, uc, nar[uc])
        r.page_break()
        r.h(3, f"4.{i}.2 Sequence Diagram")
        r.figure(img(f"SQ0{i}.png"), f"Sequence Diagram {uc} {UC_NAMES[uc]}", max_h=22)
        r.h(3, f"4.{i}.3 Collaboration Diagram")
        r.figure(img(f"CL0{i}.png"), f"Collaboration Diagram {uc} {UC_NAMES[uc]}")


def csv_rows(table):
    with open(os.path.join(DATA, f"{table}.csv"), encoding="utf-8") as f:
        return list(csv.DictReader(f))


def short(v, col):
    if v is None:
        return ""
    v = str(v)
    if col == "Password" and v.startswith("$2"):
        return v[:12] + "…"
    if re.match(r"\d{4}-\d\d-\d\d \d\d:\d\d:\d\d", v):
        v = v[:16]
    if col in ("ProfilePicture", "FilePath", "AttachmentFile") and len(v) > 22:
        v = "…" + v[-18:]
    if len(v) > 60:
        v = v[:57] + "…"
    return v if v != "" else "NULL"


SAMPLE_PICK = {"NOTIFICATION": lambda rows: [rows[i] for i in sorted({0, len(rows) // 5, len(rows) // 3, len(rows) // 2, len(rows) - 2})],
               "PROJECT_DOCUMENT": lambda rows: rows[:2] + rows[-3:],
               "ENGINEER_SKILL": lambda rows: rows[:5]}


def ch5(r):
    r.h(1, "บทที่ 5 การออกแบบฐานข้อมูล", new_page=True)
    r.h(2, "5.1 ER Diagram (Peter Chen)")
    for t in T.ER_DESC:
        r.p(t)
    r.figure(img("ER_Diagram.png"), "ER Diagram แบบ Peter Chen")
    r.h(2, "5.2 ER Diagram แบบ Logical Schema", new_page=True)
    r.p("เนื่องจากมีความสัมพันธ์แบบ M:N จึงแสดง ER Diagram แบบ Logical Schema เพิ่มเติม ชื่อตารางตรงกับ CRUD Table "
        "เส้นเชื่อมลาก FK → PK ด้วยสัญลักษณ์ Crow's foot")
    r.figure(img("ER_Logical_Schema.png"), "ER Diagram แบบ Logical Schema (11 ตาราง)")
    r.h(2, "5.3 ERD ที่มี Query", new_page=True)
    r.p(T.ERDQ_DESC)
    for i, uc in enumerate(UC_NAMES, 1):
        if i > 1:
            r.page_break()
        r.figure(img(f"ERDQ0{i}.png"), f"ERD ที่มี Query {uc} {UC_NAMES[uc]}", max_h=17)
    r.h(2, "5.4 ตาราง Query")
    r.p("ตาราง Query รวมคำสั่ง SQL ทุกคำสั่งของ UC01–UC08 (69 Query) คำสั่งเดียวกับใน Use-Case Narrative With Query และไฟล์ ProjectSA_SQL.sql "
        "ค่า @ตัวแปร คือค่าที่รับจากหน้าจอหรือผู้ใช้ที่ login ซึ่งในโปรแกรมผูกค่าเป็น parameter (:ตัวแปร) ทุกครั้ง")
    qs = queries()
    rows = [(q["qid"], q["uc"], UC_NAMES[q["uc"]], q["sql"] or "(บันทึกพร้อม Query 5.2)") for q in qs]
    r.caption("Table", "ตาราง Query (Query ID, UC ID, UC Name และ SQL Command)")
    r.table(("Query ID", "UC ID", "UC Name", "SQL Command"), rows, (1.5, 1.4, 2.9, 10.2), size=12, mono_cols=(3,),
            center_cols=(0, 1), mono_size=7.5)
    r.h(2, "5.5 Table Structure (Data Dictionary)", new_page=True)
    r.p("โครงสร้างตารางทั้ง 11 ตาราง แต่ละตารางแสดงชื่อ attribute คำอธิบาย PK / FK, Not Null, ชนิดข้อมูล (ขนาด) และช่วงค่า / การตรวจสอบ "
        "ตามด้วยตัวอย่างข้อมูล 5 รายการ ตัวอย่างข้อมูลเป็นข้อมูลที่บันทึกผ่าน Use Case จริงในโปรแกรม (สำเนาฐานข้อมูลสาธิต)")
    sch = schema()
    for n, t in enumerate(TABLE_ORDER, 1):
        r.h(3, f"5.5.{n} ตาราง {t}")
        r.p(f"{D.TABLE_DESC[t]}", indent=False)
        rows = []
        for c in sch[t]:
            key = ", ".join(x for x in (("PK" if c["pk"] else ""), ("FK → " + c["fk"] if c["fk"] else "")) if x)
            desc, val = D.COL[(t, c["name"])]
            rows.append((c["name"], desc, key, "N" if c["null"] else "Y", c["type"], val))
        r.caption("Table", f"โครงสร้างตาราง {t}")
        r.table(("Attribute", "Description", "PK / FK", "Not Null", "Data Type (Size)", "Range / Validation"), rows,
                (3.0, 3.3, 2.5, 1.2, 2.4, 3.6), size=12, center_cols=(3,))
        data = csv_rows(t)
        pick = SAMPLE_PICK.get(t, lambda rr: rr[:5])(data)[:5]
        if len(pick) < 5:
            pick = data[:5]
        srows = [(c["name"],) + tuple(short(rec.get(c["name"]), c["name"]) for rec in pick) for c in sch[t]]
        r.caption("Table", f"ตัวอย่างข้อมูลตาราง {t} ({len(pick)} รายการ)")
        r.table(("Attribute", "#1", "#2", "#3", "#4", "#5"), srows, (3.0, 2.6, 2.6, 2.6, 2.6, 2.6), size=11,
                first_col_fill="F2F2F2")


def ch6(r):
    r.h(1, "บทที่ 6 UML การออกแบบเชิงวัตถุ", new_page=True)
    r.h(2, "6.1 Class Diagram")
    for t in T.CLASS_DESC:
        r.p(t)
    r.page_break()
    r.figure(img("CD01_Entity.png"), "Class Diagram – Entity Classes")
    r.page_break()
    r.figure(img("CD02_Layers_UC01-UC04.png"), "Class Diagram – Boundary / Controller / API (UC01–UC04)")
    r.page_break()
    r.figure(img("CD03_Layers_UC05-UC08.png"), "Class Diagram – Boundary / Controller / API (UC05–UC08)")
    r.h(2, "6.2 State Diagram", new_page=True)
    r.p(T.STATE_DESC)
    r.figure(img("State_Diagram.png"), "State Diagram ของ Project.Status")


def ch7(r):
    r.h(1, "บทที่ 7 การออกแบบระบบ โปรแกรม และการทดสอบ", new_page=True)
    r.h(2, "7.1 SW Development Platform")
    r.p(T.PLATFORM_DESC)
    r.figure(img("SW_Platform.png"), "Stack ของเครื่องมือที่ใช้พัฒนา (SW Development Platform)")
    r.h(2, "7.2 Site Map / UI Structure", new_page=True)
    r.p(T.SITEMAP_DESC)
    r.figure(img("Site_Map.png"), "Site Map / UI Structure (Login เป็น root แยก tree ตามผู้ใช้)", max_h=19)
    r.h(2, "7.3 User Interface Design", new_page=True)
    r.p("หน้าจอทุกหน้าตั้งชื่อตาม Use Case และตรงกับ <<Boundary>> ใน Sequence Diagram ทุกหน้ามีแถบเมนูตาม Role, กระดิ่งแจ้งเตือน, "
        "รูปโปรไฟล์ และปุ่มออกจากระบบ")
    for title, f, desc in D.UI_SCREENS:
        r.p(f"**{title}** – {desc}", indent=False).paragraph_format.keep_with_next = True
        r.figure(shot(f), f"หน้าจอ {title}", max_w=14.5, max_h=9.3)
    r.p("**หน้าหลักของแต่ละ Role** – แสดงเมนูเฉพาะ Use Case ของ Role นั้น", indent=False).paragraph_format.keep_with_next = True
    for title, f in D.HOME_SCREENS:
        r.figure(shot(f), title, max_w=12, max_h=8)
    r.h(2, "7.4 Output Design (Report)", new_page=True)
    r.p("โปรแกรมมีรายงานบนหน้าจอ 2 รายงาน ข้อมูลในรายงานมาจากการบันทึกผ่าน Use Case จริง ดังนี้")
    for k, rep in enumerate(D.REPORTS, 1):
        r.h(3, f"7.4.{k} {rep['name']}")
        r.p(f"แสดงใน {rep['uc']} – {rep['purpose']}")
        r.caption("Table", f"โครงสร้างรายงาน {rep['name']}")
        r.table(("ส่วนของรายงาน", "ข้อมูลที่แสดง", "Query / ตาราง", "ข้อมูลมาจาก"), rep["rows"], (3.0, 6.2, 4.2, 2.6), size=13)
        r.p(rep["sample"])
        r.figure(shot(rep["image"]), f"ตัวอย่าง{rep['name']}", max_w=14.5, max_h=17)
        if rep.get("image2"):
            r.figure(shot(rep["image2"]), f"ตัวอย่าง{rep['name']} เมื่อ Filter Skill = Linux", max_w=14.5, max_h=10)
    sql_statements(r)
    testing(r)


def sql_statements(r):
    r.h(2, "7.5 SQL Statement", new_page=True)
    r.p("คำสั่ง SQL ที่โปรแกรมใช้จริงทั้งหมด 87 คำสั่ง (UC01–UC08 69 คำสั่ง และ F1–F4 18 คำสั่ง) ทดสอบแล้วว่ารันได้ทุกคำสั่ง 3 วิธี")
    r.bullet("**tools/test_sql.py** – รันทุกคำสั่งตามวงจรชีวิตของโครงการ ผลลัพธ์ queries executed: 87/87, FAILED: 0")
    r.bullet("**SqlCompatibilityTest / SqlLifecycleH2Test** (JUnit) – prepare และรันทุกคำสั่งบนฐานข้อมูล H2 ที่โปรแกรมใช้จริง ผ่านทั้งหมด")
    r.bullet("**รันในโปรแกรม** – คำสั่ง SELECT ด้านล่างรันบนฐานข้อมูลสาธิตหลังบันทึกข้อมูลผ่าน UC01–UC08 ผลลัพธ์แสดงในคอลัมน์ขวา")
    qs = {(q["uc"], q["qid"]): q for q in queries()}
    res = json.load(open(os.path.join(DATA, "select_results.json"), encoding="utf-8"))
    log = open(os.path.join(DATA, "test_sql_output.txt"), encoding="utf-8").read().splitlines()
    rows = []
    for uc, qid in D.SQL_PICK:
        q = qs[(uc, qid)]
        key = f"{uc} {qid}"
        if key in res:
            p = ", ".join(f"@{k} = {v}" for k, v in res[key]["params"].items())
            out = res[key]["output"]
            lines = out.split("\n")
            if len(lines) > 7:
                lines = lines[:6] + [lines[-1]]
            result = ("ผ่าน" + (f" ({p})" if p else "") + "\n" + "\n".join(lines))
        else:
            hit = [l for l in log if l.startswith(f"PASS {uc} {qid}")]
            result = "ผ่าน – " + (hit[0][5:] if hit else "รันใน test_sql.py และ SqlLifecycleH2Test") + \
                     "\n(ข้อมูลตัวอย่างใน 5.5 บันทึกด้วยคำสั่งนี้ผ่านโปรแกรม)"
        rows.append((f"Q{qid}", uc, q["sql"], result))
    r.caption("Table", "SQL Statement ที่ใช้ในโปรแกรมและผลการทดสอบ")
    r.table(("Query No.", "UC ID", "SQL Command", "ผลการทดสอบ"), rows, (1.4, 1.2, 7.7, 5.7), size=11, mono_cols=(2, 3),
            center_cols=(0, 1), mono_size=7.5)
    junit = open(os.path.join(DATA, "junit_summary.txt"), encoding="utf-8").read().strip().splitlines()
    r.caption("Table", "ผลการทดสอบอัตโนมัติของโปรแกรม (mvnw test)")
    r.table(("Test class", "ผลลัพธ์"), [tuple(l.split(" ", 1)) for l in junit], (9.5, 6.5), size=12)


def testing(r):
    r.h(2, "7.6 การทดสอบโปรแกรมตาม Use Case", new_page=True)
    r.p("ทดสอบแต่ละ Use Case ตามวงจรชีวิตของโครงการ \"Crypton Network Infra\" ตั้งแต่สร้างโครงการจนปิดโครงการ "
        "ภาพหน้าจอแสดงทุกการบันทึกข้อมูลและผลที่แสดงใน Check Project Details (รายงาน) กรณี Alternative flow ทดสอบด้วย JUnit (ระบุชื่อ test ในตาราง)")
    for i, uc in enumerate(UC_NAMES, 1):
        r.h(3, f"7.6.{i} {uc} {UC_NAMES[uc]}", new_page=i > 1)
        rows = [(n, a, e, res, f"ดูรูป {uc} ขั้นที่ {n}" if s else "-") for n, a, e, res, s in D.TESTS[uc]]
        r.caption("Table", f"การทดสอบ {uc} {UC_NAMES[uc]}")
        t = r.table(("ขั้น", "การทดสอบ / ข้อมูลนำเข้า", "ผลที่คาดหวัง", "ผลการทดสอบ", "ภาพหน้าจอ"), rows,
                    (1.0, 5.0, 4.8, 2.2, 3.0), size=12, center_cols=(0, 3))
        # UC id / name row on top of the table
        top = t.rows[0]._tr
        new = t.add_row()
        m = new.cells[0].merge(new.cells[-1])
        rich(m.paragraphs[0], f"UC ID: {uc}    UC Name: {UC_NAMES[uc]}", size=13, bold=True)
        shade(m, "FFF2CC")
        top.addprevious(new._tr)
        for n, a, e, res, s in D.TESTS[uc]:
            if s:
                r.figure(shot(s), f"{uc} ขั้นที่ {n}: {a}", max_w=13, max_h=8.5)


def main():
    r = Rep()
    cover(r)
    front(r)
    ch1(r); ch2(r); ch3(r); ch4(r); ch5(r); ch6(r); ch7(r)
    r.save()
    print("saved", OUT)


if __name__ == "__main__":
    main()
